"""Turn a mumbai-run.sh output folder into report.xlsx.

    python3 build-report.py <run folder> <market slug>
"""
import os, re, sys, json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

run_dir, slug = sys.argv[1], sys.argv[2]
HEAD = PatternFill("solid", fgColor="1F2937"); HFONT = Font(bold=True, color="FFFFFF")
GREEN = PatternFill("solid", fgColor="D1FAE5"); AMBER = PatternFill("solid", fgColor="FEF3C7")
RED = PatternFill("solid", fgColor="FEE2E2"); GREY = PatternFill("solid", fgColor="E5E7EB")

def status(row):
    if row["rig"]: return "RIG-BOUND", GREY
    if row["success"] is None: return "NO RESULT", GREY
    if row["success"] < 50: return "FAILED", RED
    if row["success"] < 99 or (row["p95"] or 0) > 5000: return "DEGRADED", AMBER
    return "OK", GREEN

rows = []
for d in sorted(os.listdir(os.path.join(run_dir, "results"))):
    p = os.path.join(run_dir, "results", d)
    s = open(os.path.join(p, "summary.txt")).read() if os.path.exists(os.path.join(p, "summary.txt")) else ""
    m = open(os.path.join(p, "meta.txt")).read() if os.path.exists(os.path.join(p, "meta.txt")) else ""
    g = lambda k: (re.search(k + r"=([0-9.]+)", s) or [None, None])[1]
    script = (re.search(r"script=(\S+)", m) or [None, "?"])[1]
    peak = (re.search(r"PEAK_RATE=(\d+)", m) or [None, "?"])[1]
    page = (re.search(r"MARKET_SLUG=(\S+)", m) or [None, ""])[1] or (re.search(r"PAGE_PATH=(\S+)", m) or [None, ""])[1]
    rig = "sign-in" in script or ("seeded" in d and peak not in ("20", "50", "?"))
    rows.append({"run": d, "script": script, "page": page, "peak": peak, "requests": g("requests"),
                 "success": float(g("success")) if g("success") else None, "p50": g("p50"), "p95": float(g("p95")) if g("p95") else None,
                 "p99": g("p99"), "noresp": g("no_response"), "dropped": g("dropped_iterations"), "rig": rig,
                 "canary": (re.search(r'"canary":"([0-9a-f]{7})', m) or [None, "?"])[1]})

wb = Workbook(); ws = wb.active; ws.title = "Runs"
hdr = ["Run", "Script", "Page / market", "Peak req/s", "Requests", "Success %", "p50 ms", "p95 ms", "p99 ms", "No response", "Dropped", "Deployed canary", "Status"]
ws.append(hdr)
for c in ws[1]: c.fill = HEAD; c.font = HFONT
for r in rows:
    st, fill = status(r)
    ws.append([r["run"], r["script"], r["page"], r["peak"], r["requests"], r["success"], r["p50"], r["p95"], r["p99"], r["noresp"], r["dropped"], r["canary"], st])
    ws.cell(ws.max_row, len(hdr)).fill = fill
for i, w in enumerate([34, 26, 36, 10, 10, 10, 9, 9, 9, 12, 10, 14, 12], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

wt = wb.create_sheet("Write tests"); wt.append(["Test", "Attempted", "Succeeded", "Outcomes", "Wall s", "p50 ms", "Max ms"])
for c in wt[1]: c.fill = HEAD; c.font = HFONT
jl = os.path.join(run_dir, "write-tests.jsonl")
if os.path.exists(jl):
    for line in open(jl):
        try: j = json.loads(line)
        except Exception: continue
        if "test" in j: wt.append([j["test"], j["attempted"], j["succeeded"], json.dumps(j["outcomes"]), j["wall_s"], j["p50_ms"], j["max_ms"]])
for i, w in enumerate([28, 10, 10, 60, 8, 8, 8], 1): wt.column_dimensions[get_column_letter(i)].width = w

notes = wb.create_sheet("Read me")
for line in [
    f"Market: {slug}. Every row comes from results/<run>/summary.txt and meta.txt produced by tests/load/run.sh on the DGX.",
    "RIG-BOUND = the DGX's ~15 MB/s download ceiling, not the server: /sign-in (822 KB) at any rate; seeded market page above 50 req/s.",
    "DEGRADED = success under 99% or p95 over 5 s. FAILED = success under 50%.",
    "Compare with the Chess market (2026-09-15): 5 posts → 96.9% at 100 req/s; break between 200 and 300 req/s; deep links 100% under 0.5 s; db:error windows after most runs, recovering in 10 s–2 min.",
    "Write bursts: named refusals (opposite_side_held, error_rate_limit_exceeded, error_bet_serialization_exhausted) are correct behaviour; empty-body HTTP 500s coincided with a deploy last time — check the canary before/after in log.md.",
]:
    notes.append([line])
notes.column_dimensions["A"].width = 140
wb.save(os.path.join(run_dir, "report.xlsx")); print("report.xlsx written")
